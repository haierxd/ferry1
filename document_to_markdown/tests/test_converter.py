import csv
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document


SCRIPT = Path(__file__).resolve().parents[1] / 'convert.py'


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / '输入'
        self.source.mkdir()
        self.output = self.root / '输出'

    def run_cli(self, *args):
        return subprocess.run(
            [sys.executable, str(SCRIPT), '--input', str(self.source),
             '--output', str(self.output), *args],
            capture_output=True, encoding='utf-8', errors='replace')

    def make_word(self, name='样例.docx'):
        path = self.source / name
        path.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        doc.add_heading('治理规则', 0)
        doc.add_heading('采购权限', 1)
        doc.add_paragraph('恰好20000元属于部门负责人权限。')
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '版本'
        table.cell(0, 1).text = '金额'
        table.cell(1, 0).text = 'V2.0'
        table.cell(1, 1).text = '20000元'
        doc.add_paragraph('表后正文')
        doc.save(path)
        return path

    def test_word_order_and_table(self):
        self.make_word('子目录/样例.docx')
        result = self.run_cli()
        self.assertEqual(result.returncode, 0, result.stderr)
        target = self.output / '子目录/样例.docx/document.md'
        text = target.read_text(encoding='utf-8')
        self.assertIn('# 治理规则', text)
        self.assertIn('## 采购权限', text)
        self.assertIn('| V2.0 | 20000元 |', text)
        self.assertLess(text.index('20000元属于'), text.index('| V2.0'))
        self.assertLess(text.index('| V2.0'), text.index('表后正文'))

    @unittest.skipUnless(os.name == 'nt', 'Windows PowerShell entry point')
    def test_powershell_entrypoint(self):
        if not (SCRIPT.parent / '.venv/Scripts/python.exe').exists():
            self.skipTest('请先运行 run.ps1 -InstallOnly 创建环境')
        self.make_word()
        result = subprocess.run(
            ['powershell.exe', '-NoProfile', '-ExecutionPolicy', 'Bypass',
             '-File', str(SCRIPT.parent / 'run.ps1'), '-InputPath', str(self.source),
             '-OutputPath', str(self.output)], capture_output=True,
            encoding='utf-8', errors='replace', timeout=60)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertNotIn('Traceback', result.stderr)
        self.assertNotIn('Requirement already satisfied', result.stdout)
        self.assertTrue((self.output / '样例.docx/document.md').exists())

    def test_skip_only_when_source_and_outputs_unchanged(self):
        self.make_word()
        self.assertEqual(self.run_cli().returncode, 0)
        self.assertEqual(self.run_cli().returncode, 0)
        with (self.output / 'conversion_report.csv').open(encoding='utf-8-sig') as stream:
            report = list(csv.DictReader(stream))
        self.assertEqual(report[0]['status'], 'skipped')
        md = self.output / '样例.docx/document.md'
        md.write_text('damaged', encoding='utf-8')
        self.assertEqual(self.run_cli().returncode, 0)
        self.assertIn('治理规则', md.read_text(encoding='utf-8'))

    def test_failure_does_not_stop_other_files(self):
        self.make_word()
        (self.source / '损坏.docx').write_bytes(b'not a word document')
        result = self.run_cli()
        self.assertEqual(result.returncode, 1)
        with (self.output / 'conversion_report.csv').open(encoding='utf-8-sig') as stream:
            report = list(csv.DictReader(stream))
        self.assertEqual({r['status'] for r in report}, {'success', 'failed'})
        self.assertTrue((self.output / '样例.docx/document.md').exists())

    def test_reject_output_containing_input(self):
        result = self.run_cli('--output', str(self.root))
        self.assertEqual(result.returncode, 2)
        self.assertFalse((self.root / 'conversion_report.csv').exists())

    def test_real_pdf_preserves_table_and_page_mapping(self):
        sample = SCRIPT.parents[1] / '甲方资料/02_PDF合作案例/01_离散制造工厂设备联网.pdf'
        result = subprocess.run([sys.executable, str(SCRIPT), '--input', str(sample),
                                 '--output', str(self.output)], capture_output=True,
                                encoding='utf-8', errors='replace')
        self.assertEqual(result.returncode, 0, result.stderr)
        folder = self.output / (sample.name)
        text = (folder / 'document.md').read_text(encoding='utf-8')
        self.assertIn('LC-GW-400', text)
        self.assertIn('99.6', text)
        self.assertIn('| PRD-01 |', text)
        self.assertNotIn('合成教学案例', text)
        self.assertIn('标准规格。', text)
        self.assertIn('工程单位。', text)
        meta = json.loads((folder / 'metadata.json').read_text(encoding='utf-8'))
        self.assertEqual(meta['pages'], 3)
        self.assertEqual(len(meta['page_map']), 3)
        for page in meta['page_map']:
            self.assertEqual(text.splitlines()[page['markdown_line'] - 1], page['marker'])

    def test_ignore_office_temporary_files(self):
        self.make_word()
        (self.source / '.~_样例.docx').write_bytes(b'office lock')
        (self.source / '~$样例.docx').write_bytes(b'office lock')
        result = self.run_cli()
        self.assertEqual(result.returncode, 0, result.stdout)
        with (self.output / 'conversion_report.csv').open(encoding='utf-8-sig') as stream:
            self.assertEqual(len(list(csv.DictReader(stream))), 1)

    def test_merged_cells_and_embedded_image(self):
        from PIL import Image
        doc = Document()
        doc.add_heading('合并表格测试', 0)
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).merge(t.cell(1, 0)).text = '跨两行'
        t.cell(0, 1).text = '上'
        t.cell(1, 1).text = '下'
        image = self.root / 'image.png'
        Image.new('RGB', (10, 10), color='red').save(image)
        doc.add_picture(str(image))
        doc.save(self.source / '图片.docx')
        result = self.run_cli()
        self.assertEqual(result.returncode, 1, result.stderr)
        folder = self.output / '图片.docx'
        text = (folder / 'document.md').read_text(encoding='utf-8')
        self.assertIn('rowspan="2"', text)
        self.assertEqual(text.count('跨两行'), 1)
        self.assertIn('assets/image_001.png', text)
        self.assertTrue((folder / 'assets/image_001.png').is_file())

    def test_blank_pdf_requires_review(self):
        objects = [b'<< /Type /Catalog /Pages 2 0 R >>',
                   b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
                   b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << >> >>']
        data = b'%PDF-1.4\n'
        offsets = [0]
        for i, obj in enumerate(objects, 1):
            offsets.append(len(data))
            data += f'{i} 0 obj\n'.encode() + obj + b'\nendobj\n'
        start = len(data)
        data += b'xref\n0 4\n0000000000 65535 f \n'
        data += b''.join(f'{n:010d} 00000 n \n'.encode() for n in offsets[1:])
        data += f'trailer\n<< /Size 4 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF'.encode()
        (self.source / '空白.pdf').write_bytes(data)
        result = self.run_cli()
        self.assertEqual(result.returncode, 1, result.stderr)
        meta = json.loads((self.output / '空白.pdf/metadata.json').read_text(encoding='utf-8'))
        self.assertEqual(meta['status'], 'needs_review')
        self.assertIn('OCR', ''.join(meta['warnings']))

    def test_output_link_cannot_overwrite_source(self):
        source = self.make_word()
        original = source.read_bytes()
        target = self.output / '样例.docx'
        target.mkdir(parents=True)
        try:
            (target / 'document.md').symlink_to(source)
        except OSError:
            if os.name != 'nt':
                self.skipTest('此系统未允许创建符号链接')
            junction = subprocess.run(['cmd.exe', '/c', 'mklink', '/J',
                                       str(target / 'assets'), str(self.source)],
                                      capture_output=True)
            if junction.returncode:
                self.skipTest('此系统未允许创建符号链接或目录联接')
        result = self.run_cli('--force')
        self.assertEqual(result.returncode, 1)
        self.assertEqual(source.read_bytes(), original)

    def test_old_dependencies_are_not_accepted(self):
        sys.path.insert(0, str(SCRIPT.parent))
        self.addCleanup(sys.path.remove, str(SCRIPT.parent))
        import check_environment
        with patch.object(check_environment, 'version', return_value='0.8.11'):
            self.assertFalse(check_environment.ready('basic'))


if __name__ == '__main__':
    unittest.main()
