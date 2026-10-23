"""Local DOCX/PDF to Markdown. Document text is data; no macros or links execute."""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
import csv
import hashlib
import html
from importlib.metadata import version
import json
import os
from pathlib import Path
import re
import shutil
import statistics
import sys
import tempfile
from urllib.parse import quote

VERSION = '1.0.2'
BASE = Path(__file__).resolve().parents[1]


def digest(path):
    h = hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def reject_links(path):
    """Do not follow pre-existing output links/junctions while reading or writing."""
    def check(node):
        if node.is_symlink() or getattr(node, 'is_junction', lambda: False)():
            raise ValueError(f'输出路径包含符号链接或目录联接，已拒绝写入：{node}')
    for node in [path, *path.parents]:
        check(node)
    if path.is_dir():
        for directory, dirs, files in os.walk(path, followlinks=False):
            for name in dirs + files:
                check(Path(directory) / name)


def escape(text):
    # Prevent source text from becoming HTML, links or Markdown commands.
    text = str(text or '').replace('\r', '').replace('\x00', '')
    text = html.escape(text, quote=False)
    return re.sub(r'([\\`*_{}\[\]<>#|])', r'\\\1', text)


def compact(text):
    # Join only between Chinese characters, avoiding damage to English words.
    return re.sub(r'(?<=[\u3400-\u9fff]) +(?=[\u3400-\u9fff])', '', text)


def markdown_table(rows):
    if not rows:
        return ''
    width = max(map(len, rows))
    data = [[escape(c).replace('\n', '<br>') for c in row] + [''] * (width - len(row)) for row in rows]
    line = lambda cells: '| ' + ' | '.join(cells) + ' |'
    return '\n'.join([line(data[0]), line(['---'] * width)] + [line(r) for r in data[1:]])


def convert_docx(source, folder):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(source)
    warnings = []
    blocks = []
    image_names = {}
    counters = {}
    has_title = any(p.style and p.style.name == 'Title' for p in doc.paragraphs)

    def run_text(run):
        text = escape(run.text)
        if text.strip():
            if run.bold:
                text = '**' + text + '**'
            if run.italic:
                text = '*' + text + '*'
        for blip in run._r.xpath('.//a:blip'):
            rid = blip.get(qn('r:embed'))
            if not rid:
                warnings.append('外部链接图片未下载，请人工核对。')
                continue
            part = doc.part.related_parts[rid]
            if rid not in image_names:
                assets = folder / 'assets'
                assets.mkdir(exist_ok=True)
                name = f'image_{len(image_names) + 1:03d}{Path(part.partname).suffix}'
                (assets / name).write_bytes(part.blob)
                image_names[rid] = name
            text += f' ![图片](assets/{quote(image_names[rid])}) '
        return text

    def paragraph(p):
        from docx.text.hyperlink import Hyperlink
        text_parts = []
        for item in p.iter_inner_content():
            if isinstance(item, Hyperlink):
                label = ''.join(run_text(r) for r in item.runs)
                url = item.url
                # Preserve unsafe/unsupported schemes as text, never active links.
                if re.match(r'^(https?://|mailto:)', url, re.I):
                    text_parts.append(f'[{label}]({quote(url, safe=":/?=&%#@+~")})')
                else:
                    text_parts.append(label + (f' ({escape(url)})' if url else ''))
            else:
                text_parts.append(run_text(item))
        text = ''.join(text_parts).strip()
        if not text:
            return ''
        style = p.style.name if p.style else ''
        if style == 'Title':
            return '# ' + text
        match = re.match(r'Heading\s*(\d+)', style, re.I)
        if match:
            level = min(6, int(match.group(1)) + int(has_title))
            return '#' * level + ' ' + text
        prop = p._p.pPr
        num = prop.numPr if prop is not None else None
        if num is not None and num.numId is not None and num.numId.val != 0:
            level = num.ilvl.val if num.ilvl is not None else 0
            key = (num.numId.val, level)
            try:
                numbering = doc.part.numbering_part.element
                abstract_id = numbering.xpath(f'./w:num[@w:numId="{key[0]}"]/w:abstractNumId')[0].val
                levels = numbering.xpath(f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]/w:lvl[@w:ilvl="{level}"]')
                fmt = levels[0].find(qn('w:numFmt')).get(qn('w:val'))
                start_el = levels[0].find(qn('w:start'))
                start = int(start_el.get(qn('w:val'))) if start_el is not None else 1
                counters[key] = counters.get(key, start - 1) + 1
                marker = '-' if fmt == 'bullet' else f'{counters[key]}.'
                if fmt not in ('bullet', 'decimal'):
                    warnings.append('自动编号已规范化为数字列表，原编号格式请复核。')
                return '  ' * level + marker + ' ' + text
            except (IndexError, AttributeError, KeyError, ValueError):
                warnings.append('自动编号无法完整解析，已转换为普通列表。')
                return '  ' * level + '- ' + text
        if style.startswith('List Bullet'):
            return '- ' + text
        if style.startswith('List Number'):
            counters[style] = counters.get(style, 0) + 1
            return str(counters[style]) + '. ' + text
        return text.replace('\n', '  \n')

    def table(t):
        merged = bool(t._tbl.xpath('.//w:gridSpan | .//w:vMerge'))
        if not merged:
            return markdown_table([[c.text for c in r.cells] for r in t.rows])
        warnings.append('合并单元格表格已输出 HTML；请使用支持 HTML 的 Markdown 阅读器。')
        # Work with physical XML cells to preserve horizontal/vertical spans.
        result = ['<table>']
        rows = t._tbl.findall(qn('w:tr'))
        for index, row in enumerate(rows):
            result.append('<tr>')
            column = 0
            for cell in row.findall(qn('w:tc')):
                props = cell.find(qn('w:tcPr'))
                span_el = props.find(qn('w:gridSpan')) if props is not None else None
                span = int(span_el.get(qn('w:val'))) if span_el is not None else 1
                merge = props.find(qn('w:vMerge')) if props is not None else None
                if merge is not None and merge.get(qn('w:val')) != 'restart':
                    column += span
                    continue
                rowspan = 1
                if merge is not None:
                    for later in rows[index + 1:]:
                        position = 0
                        continuation = False
                        for other in later.findall(qn('w:tc')):
                            op = other.find(qn('w:tcPr'))
                            gs = op.find(qn('w:gridSpan')) if op is not None else None
                            vm = op.find(qn('w:vMerge')) if op is not None else None
                            if position == column:
                                continuation = vm is not None and vm.get(qn('w:val')) != 'restart'
                                break
                            position += int(gs.get(qn('w:val'))) if gs is not None else 1
                        if not continuation:
                            break
                        rowspan += 1
                paras = cell.findall(qn('w:p'))
                content = '<br>'.join(html.escape(''.join(p.xpath('.//w:t/text()'))) for p in paras)
                result.append(f'<td colspan="{span}" rowspan="{rowspan}">{content}</td>')
                column += span
            result.append('</tr>')
        return '\n'.join(result + ['</table>'])

    for child in doc.element.body:
        if child.tag == qn('w:p'):
            blocks.append(paragraph(Paragraph(child, doc)))
        elif child.tag == qn('w:tbl'):
            blocks.append(table(Table(child, doc)))
        elif child.tag != qn('w:sectPr'):
            warnings.append('存在内容控件或其他特殊正文结构，请人工复核。')
    if doc.element.xpath('.//w:txbxContent | .//w:ins | .//w:del'):
        warnings.append('检测到文本框或修订记录，基础模式不保证完整保留，请人工复核。')
    if doc.element.xpath('.//m:oMath | .//w:footnoteReference | .//w:endnoteReference'):
        warnings.append('检测到公式或脚注尾注，请使用高级模式并人工复核。')
    for t in doc.tables:
        if t._tbl.xpath('.//w:tbl | .//w:drawing'):
            warnings.append('表格包含嵌套表格或图片，请人工复核。')
    text = '\n\n'.join(b for b in blocks if b)
    if not text.strip():
        warnings.append('没有提取到正文。')
    return text + '\n', {'pages': None, 'page_map': [], 'warnings': warnings,
                         'tables': len(doc.tables), 'images': len(image_names)}


def pdf_lines(page):
    # Text lines retain geometry and font sizes for table exclusion and headings.
    words = page.extract_words(extra_attrs=['size'], x_tolerance=2, y_tolerance=3)
    lines = []
    for word in sorted(words, key=lambda w: (round(w['top'] / 3), w['x0'])):
        if not lines or abs(word['top'] - lines[-1]['top']) > 3:
            lines.append({'top': word['top'], 'bottom': word['bottom'], 'words': []})
        lines[-1]['words'].append(word)
        lines[-1]['bottom'] = max(lines[-1]['bottom'], word['bottom'])
    for line in lines:
        ordered = sorted(line['words'], key=lambda w: w['x0'])
        line['text'] = compact(' '.join(w['text'] for w in ordered))
        line['size'] = statistics.median(w['size'] for w in ordered)
    return lines


def convert_pdf(source, folder, keep_headers=False):
    import pdfplumber
    warnings, sections, mapping = [], [], []
    total_tables = total_images = 0
    with pdfplumber.open(source) as pdf:
        # Only remove repeated text inside narrow top/bottom margins.
        repeated = Counter()
        margin_key = lambda text: re.sub(r'\s+\d+\s*$', '', text)
        all_lines = []
        for page in pdf.pages:
            lines = pdf_lines(page)
            all_lines.append(lines)
            repeated.update(set(margin_key(l['text']) for l in lines if l['top'] < 42 or l['bottom'] > page.height - 36))
        for page_num, (page, lines) in enumerate(zip(pdf.pages, all_lines), 1):
            marker = f'<!-- source_page: {page_num} -->'
            page_blocks = [marker]
            # Require multiple rows and columns to avoid treating paragraph boxes as tables.
            tables = [t for t in page.find_tables() if len(t.rows) >= 2 and len(t.columns) >= 2]
            events = []
            for t in tables:
                events.append((t.bbox[1], 'table', markdown_table(t.extract())))
                total_tables += 1
            body_size = statistics.median([c['size'] for c in page.chars]) if page.chars else 10
            buffer = []
            previous_bottom = None

            def flush():
                if buffer:
                    page_blocks.append(''.join(buffer))
                    buffer.clear()

            for line in lines:
                text = line['text']
                marginal = line['top'] < 42 or line['bottom'] > page.height - 36
                if not keep_headers and marginal:
                    if re.fullmatch(r'(?:第\s*)?\d+\s*(?:页)?', text) or repeated[margin_key(text)] >= 2:
                        continue
                # Remove only words inside table boxes (retain any nearby caption).
                words = [w for w in line['words'] if not any(
                    t.bbox[0] <= (w['x0'] + w['x1']) / 2 <= t.bbox[2]
                    and t.bbox[1] <= (w['top'] + w['bottom']) / 2 <= t.bbox[3] for t in tables)]
                if not words:
                    continue
                text = compact(' '.join(w['text'] for w in sorted(words, key=lambda w: w['x0'])))
                events.append((line['top'], 'line', {**line, 'text': text}))
            # Images are page crops, including any overlaid labels; not reconstructed originals.
            for image_num, item in enumerate(page.images, 1):
                bbox = (max(0, item['x0']), max(0, item['top']),
                        min(page.width, item['x1']), min(page.height, item['bottom']))
                if bbox[2] <= bbox[0] or bbox[3] <= bbox[1]:
                    continue
                assets = folder / 'assets'
                assets.mkdir(exist_ok=True)
                name = f'page_{page_num:04d}_image_{image_num:03d}.png'
                try:
                    page.crop(bbox).to_image(resolution=144).save(assets / name)
                    events.append((bbox[1], 'image', f'![第{page_num}页图片](assets/{name})'))
                    total_images += 1
                except Exception as exc:
                    warnings.append(f'第{page_num}页图片导出失败：{exc}')
            visible = ''.join(l['text'] for l in lines if 42 <= l['top'] and l['bottom'] <= page.height - 36)
            if len(visible.strip()) < 20:
                warnings.append(f'第{page_num}页正文文字过少，可能是扫描页或空白页，需要 OCR 或人工核对。')
            for top, kind, item in sorted(events, key=lambda e: e[0]):
                if kind != 'line':
                    flush()
                    page_blocks.append(item)
                    previous_bottom = None
                    continue
                text = escape(item['text'])
                heading = item['size'] >= body_size * 1.18 or bool(re.match(r'^[一二三四五六七八九十]+[、\s]', item['text']))
                if heading and len(text) < 110:
                    flush()
                    text = re.sub(r'^([一二三四五六七八九十]+)(?=[^一二三四五六七八九十\s])', r'\1 ', text)
                    prefix = '# ' if page_num == 1 and item['size'] >= body_size * 1.6 else '## '
                    page_blocks.append(prefix + text)
                else:
                    if previous_bottom is not None and top - previous_bottom > body_size * 1.1:
                        flush()
                    if buffer and re.search(r'[A-Za-z0-9]$', buffer[-1]) and re.match(r'^[A-Za-z0-9]', text):
                        buffer.append(' ')
                    buffer.append(text)
                previous_bottom = item['bottom']
            flush()
            start_line = sum(s.count('\n') + 2 for s in sections) + 1
            section = '\n\n'.join(page_blocks)
            sections.append(section)
            mapping.append({'page': page_num, 'marker': marker, 'markdown_line': start_line})
        return '\n\n'.join(sections) + '\n', {
            'pages': len(pdf.pages), 'page_map': mapping, 'warnings': warnings,
            'tables': total_tables, 'images': total_images}


def convert_advanced(source, folder):
    try:
        from docling.document_converter import DocumentConverter
        from docling_core.types.doc import ImageRefMode
    except ImportError as exc:
        raise RuntimeError('高级模式需要 Docling：python -m pip install -r requirements-advanced.txt') from exc
    result = DocumentConverter().convert(str(source))
    if 'success' not in str(result.status).lower() or 'partial' in str(result.status).lower():
        raise RuntimeError(f'Docling 未完整成功：{result.status}; {result.errors}')
    result.document.save_as_markdown(folder / 'document.md', image_mode=ImageRefMode.REFERENCED)
    text = (folder / 'document.md').read_text(encoding='utf-8')
    # Keep the structured document for page/bounding-box provenance.
    result.document.save_as_json(folder / 'structure.json')
    return text, {'pages': len(result.document.pages), 'page_map': [],
                  'warnings': ['高级模式的页码和坐标请查 structure.json；中文 OCR 与复杂表格需要人工抽查。'],
                  'tables': len(result.document.tables), 'images': len(result.document.pictures)}


def process(source, target, args):
    reject_links(target)
    source_hash = digest(source)
    options = {'engine': args.engine, 'keep_pdf_headers': args.keep_pdf_headers}
    meta_file = target / 'metadata.json'
    if not args.force and meta_file.exists():
        try:
            meta = json.loads(meta_file.read_text(encoding='utf-8'))
            files = meta.get('output_sha256', {})
            valid = files and all((target / f).is_file() and digest(target / f) == h for f, h in files.items())
            if (meta.get('source_sha256') == source_hash and meta.get('tool_version') == VERSION
                    and meta.get('options') == options and valid):
                return {'status': 'skipped', 'warnings': '；'.join(meta.get('warnings', [])), 'error': ''}
        except (OSError, ValueError, TypeError):
            pass
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.convert-', dir=target.parent) as temp:
        folder = Path(temp)
        if args.engine == 'docling':
            text, detail = convert_advanced(source, folder)
        elif source.suffix.lower() == '.docx':
            text, detail = convert_docx(source, folder)
        else:
            text, detail = convert_pdf(source, folder, args.keep_pdf_headers)
        if not text.strip():
            detail['warnings'].append('输出正文为空，需要人工核对。')
        (folder / 'document.md').write_text(text, encoding='utf-8')
        warnings = list(dict.fromkeys(detail.pop('warnings')))
        status = 'needs_review' if warnings else 'success'
        dependencies = {}
        for name in (['docling', 'docling-core'] if args.engine == 'docling' else ['python-docx', 'pdfplumber']):
            dependencies[name] = version(name)
        meta = {'source': str(source), 'source_sha256': source_hash,
                'tool_version': VERSION, 'dependencies': dependencies,
                'converted_at': datetime.now(timezone.utc).isoformat(),
                'options': options, 'status': status, 'warnings': warnings, **detail,
                'output_sha256': {f.relative_to(folder).as_posix(): digest(f)
                                  for f in folder.rglob('*') if f.is_file()}}
        (folder / 'metadata.json').write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding='utf-8')
        # Replace generated files only; retain untracked user files in the result directory.
        target.mkdir(exist_ok=True)
        for item in folder.iterdir():
            if item.is_dir():
                shutil.copytree(item, target / item.name, dirs_exist_ok=True)
            else:
                shutil.copy2(item, target / item.name)
    return {'status': status, 'warnings': '；'.join(warnings), 'error': ''}


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8', errors='replace')
    parser = argparse.ArgumentParser(description='本地批量转换 DOCX/PDF 为 Markdown；默认不联网、不调用大模型。')
    parser.add_argument('--input', type=Path, default=BASE / '甲方资料', help='文件或目录，递归扫描')
    parser.add_argument('--output', type=Path, default=BASE / 'markdown_output', help='输出目录')
    parser.add_argument('--engine', choices=['basic', 'docling'], default='basic')
    parser.add_argument('--force', action='store_true', help='忽略哈希缓存，重新生成')
    parser.add_argument('--keep-pdf-headers', action='store_true', help='保留 PDF 页眉页码')
    args = parser.parse_args(argv)
    source, output = args.input.resolve(), args.output.resolve()
    if not source.exists():
        parser.error(f'输入不存在：{source}')
    root = source if source.is_dir() else source.parent
    if output == root or output in root.parents or output == source:
        parser.error('输出目录不能等于输入目录或包含输入目录。')
    try:
        reject_links(output / 'conversion_report.csv')
    except ValueError as exc:
        parser.error(str(exc))
    candidates = [source] if source.is_file() else sorted(source.rglob('*'))
    files = [p for p in candidates if p.is_file() and p.suffix.lower() in ('.docx', '.pdf')
             and not p.name.startswith(('~$', '.~', '~WRL')) and output not in p.parents]
    unsupported = [p for p in candidates if p.is_file() and p.suffix.lower() == '.doc'
                   and output not in p.parents]
    if not files and not unsupported:
        parser.error('没有找到 .docx 或 .pdf 文件；旧版 .doc 请先另存为 .docx。')
    output.mkdir(parents=True, exist_ok=True)
    report = []
    for index, path in enumerate(files + unsupported, 1):
        # Retain the extension in the folder name to prevent same-stem collisions.
        target = output / path.relative_to(root) / 'document.md'
        print(f'[{index}/{len(files) + len(unsupported)}] {path.name}', flush=True)
        try:
            if path.suffix.lower() == '.doc':
                raise ValueError('旧版 .doc 不支持，请用 Word 另存为 .docx 后重试。')
            result = process(path, target.parent, args)
        except Exception as exc:
            result = {'status': 'failed', 'warnings': '', 'error': f'{type(exc).__name__}: {exc}'}
        report.append({'source': str(path), 'output': str(target), **result})
        print('  ' + result['status'] + ('：' + result['error'] if result['error'] else ''), flush=True)
    with (output / 'conversion_report.csv').open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=['source', 'output', 'status', 'warnings', 'error'])
        writer.writeheader()
        writer.writerows(report)
    counts = Counter(r['status'] for r in report)
    print('完成：' + json.dumps(dict(counts), ensure_ascii=False))
    print(f'报告：{output / "conversion_report.csv"}')
    return 1 if counts['failed'] or counts['needs_review'] else 0


if __name__ == '__main__':
    raise SystemExit(main())
